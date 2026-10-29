"""Regenerate fictional fixtures: pip install python-docx reportlab. No personal career data."""
from pathlib import Path
from docx import Document
from reportlab.pdfgen import canvas
from reportlab.lib.utils import simpleSplit
ROOT = Path(__file__).resolve().parents[1] / 'public/sample-resumes'
ROOT.mkdir(parents=True, exist_ok=True)
keys = ['requirements','process','uat','roadmap','discovery','backlog','planning','stakeholders','delivery','sql','pipelines','warehouse']
ba = [
 ('Demonstrated','Led discovery workshops and wrote checkout acceptance criteria at Fictional Retail, 2022-2025.'),
 ('Demonstrated','Mapped order fulfilment processes and agreed improvements with operations at Fictional Retail.'),
 ('Demonstrated','Coordinated UAT with business users and triaged release defects at Fictional Retail.'),
 ('Adjacent','Proposed quarterly roadmap options; the product manager owned final roadmap decisions.'),
 ('Adjacent','Supported customer interviews and synthesized findings with the product manager.'),
 ('Demonstrated','Prioritized the checkout delivery backlog within agreed product objectives.'),
 ('Demonstrated','Owned release plans, dependencies and risk tracking for a fictional checkout migration.'),
 ('Demonstrated','Aligned engineers and operations stakeholders on release scope and changes.'),
 ('Demonstrated','Owned delivery reporting and release coordination for the checkout migration.'),
 ('Limited','Used basic SELECT queries; has not written complex production SQL transformations.'),
 ('No evidence','No production data pipeline work is included in this fictional career record.'),
 ('No evidence','No warehouse architecture work is included in this fictional career record.')]
engineer=[('Limited','Observed business analysis workshops but did not lead this work.')]*3 + [('Unknown','Product ownership scope is not described in this fictional resume.')]*3 + [('Adjacent','Coordinated technical tasks within a data platform team.')]*3 + [
 ('Demonstrated','Owned complex production SQL transformations and performance tuning at Fictional Data, 2021-2025.'),
 ('Demonstrated','Built monitored data pipelines and handled production recovery at Fictional Data.'),
 ('Demonstrated','Designed warehouse models and quality controls for subscription analytics at Fictional Data.')]
incomplete=[('Unknown','This short fictional profile does not describe comparable project evidence.')]*12
for name, person, summary, evidence in [
 ('business-analyst','Alex Morgan','Fictional digital business analyst with ecommerce delivery experience.',ba),
 ('data-engineer','Sam Chen','Fictional analytics engineer focused on production data infrastructure.',engineer),
 ('incomplete-profile','Jordan Reed','Fictional incomplete profile used to demonstrate investigation rather than guessing.',incomplete)]:
 paragraphs=[person+' - FICTIONAL SAMPLE RESUME',summary,
 'Prototype evidence notes: explicit assessments for a controlled rules demonstration. These are fictional claims, not verified facts. This format is not a general-purpose resume understanding system.']
 paragraphs += [f'Evidence [{key}]: {level} | {quote}' for key,(level,quote) in zip(keys,evidence)]
 paragraphs += ['French eligibility: Unknown']
 d=Document();d.add_heading(paragraphs[0],0)
 for text in paragraphs[1:]:d.add_paragraph(text)
 d.save(ROOT / (name+'.docx'))
 c=canvas.Canvas(str(ROOT / (name+'.pdf')));c.setTitle(person+' - Fictional sample resume');y=800
 for paragraph in paragraphs:
  for line in simpleSplit(paragraph,'Helvetica',10,480):
   if y<55:c.showPage();y=800
   c.setFont('Helvetica',10);c.drawString(55,y,line);y-=14
  y-=10
 c.save()

# Ordinary resumes with no special evidence tags or preassigned assessment levels.
for name, person, entries in [
 ('ordinary-business-analyst','Alex Morgan',[
  'Digital Business Analyst | Fictional Retail | 2022-2025',
  'Led discovery workshops and wrote checkout requirements and acceptance criteria for a retail ordering platform.',
  'Mapped fulfilment processes with operations and agreed measurable workflow improvements.',
  'Coordinated UAT with business users and triaged release defects against the agreed acceptance criteria.',
  'Supported customer interviews and proposed roadmap options; the product manager owned final decisions.',
  'Prioritized the checkout backlog within agreed product objectives.',
  'Owned release plans, risk tracking and delivery reporting; aligned operations stakeholders and engineers on scope.',
  'Used basic SQL SELECT queries. No experience owning production data pipelines or warehouse architecture.',
  'Education: BA in Business, Fictional College. Languages: English.']),
 ('ordinary-data-engineer','Sam Chen',[
  'Analytics Engineer | Fictional Data | 2021-2025',
  'Owned complex production SQL transformations and performance tuning for subscription analytics.',
  'Built and monitored data pipelines; handled recovery and data quality incidents.',
  'Designed warehouse models and quality controls for finance and product reporting.',
  'Supported business stakeholders in interpreting data; product managers owned the roadmap and discovery.',
  'Education: BSc in Computing, Fictional College. Languages: English.'])]:
 paragraphs=[person+' - FICTIONAL SAMPLE RESUME']+entries
 d=Document();d.add_heading(paragraphs[0],0)
 for text in paragraphs[1:]:d.add_paragraph(text)
 d.save(ROOT/(name+'.docx'))
 c=canvas.Canvas(str(ROOT/(name+'.pdf')));y=800
 for paragraph in paragraphs:
  for line in simpleSplit(paragraph,'Helvetica',11,480):
   if y<55:c.showPage();y=800
   c.setFont('Helvetica',11);c.drawString(55,y,line);y-=16
  y-=12
 c.save()
